from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart
#B2:C11 까지의 데이터를 차트로 생성
bar_value =  Reference(ws, min_row = 2 , max_row = 11, min_col= 2, max_col=3) # 어떤 범위의 데이터를 쓸지 설정
bar_chart = BarChart() # 차트 종류 설정 (bar, line, pie)
bar_chart.add_data(bar_value) # 바 차트에 미리 설정한 데이터 추가
ws.add_chart(bar_chart, "E1") # 워크시트에 설정된 위치로 차트 추가

# B1:C11 까지의 데이터
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3) # 어떤 범위의 데이터를 쓸지 설정
line_chart = LineChart() # 차트 종류 설정
line_chart.add_data(line_value, titles_from_data = True) # 계열을 > 영어, 수학 (제목에서 가져옴)
line_chart.title = "성적표" # 차트 제목 부여
line_chart.style = 10 # 미리 정의된 스타일ㅇ을 적용, 사용자가 개별 지정도 가능
line_chart.y_axis.title = "점수" # Y축의 제목
line_chart.x_axis.title = "번호" # X축의 제목
ws.add_chart(line_chart, "E1") # 워크시트에 설정된 위치로 차트 추가

wb.save("sample_chart.xlsx")
